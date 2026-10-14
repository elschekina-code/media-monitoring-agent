# -*- coding: utf-8 -*-
"""
Génère le classeur Excel (.xlsx) détaillé des retombées presse (MEDIA MONITORING AGENT V1.1).

Usage :
    python scripts/generer_excel.py --cp CP-20261010-001 [--out ...]

Onglets : EXECUTIVE SUMMARY, MEDIA COVERAGE, GEOGRAPHIC ANALYSIS,
PUBLICATION TIMELINE, VERIFICATION LOG.
Chaque publication : identifiant, média, pays, date, titre, statut, URL cliquable.
Marqué « BROUILLON — EN ATTENTE DE VALIDATION ». Aucun envoi automatique.
"""
import argparse
import csv
import json
import pathlib
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

BASE = pathlib.Path(__file__).resolve().parent.parent


def lire_reprises(cp):
    f = BASE / "reprises.csv"
    rows = list(csv.DictReader(f.open(encoding="utf-8"), delimiter=";"))
    return [r for r in rows if r.get("communique_ref") == cp and r.get("est_test") == "non"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cp", default="CP-20261010-001")
    ap.add_argument("--config", default=str(BASE / "config" / "charte.json"))
    ap.add_argument("--out", default="")
    ap.add_argument("--titre", default="")
    args = ap.parse_args()

    cfg = json.loads(pathlib.Path(args.config).read_text(encoding="utf-8"))
    coul = cfg.get("couleurs", {})
    PRIM = coul.get("primaire", "1F3A5F").lstrip("#")
    ACC = coul.get("accent", "C9A227").lstrip("#")
    filigrane = cfg.get("filigrane", "BROUILLON — EN ATTENTE DE VALIDATION")

    rows = lire_reprises(args.cp)
    conf = [r for r in rows if r.get("classification") == "confirmee"]
    prob = [r for r in rows if r.get("classification") == "probable"]
    nv = [r for r in rows if r.get("classification") == "non_verifiable"]
    uniques = conf + prob

    titre = args.titre or (rows[0]["titre"] if rows else "")
    nb_conf = len(conf)
    nb_prob = len(prob)
    nb_nv = len(nv)
    medias_uniques = len({r["media"] for r in uniques})
    pays_counts = Counter(r["pays"] for r in uniques if r.get("pays") and r["pays"] not in ("—", ""))
    pays_nd = [r for r in uniques if r.get("pays") in ("Non déterminé", "")]
    type_counts = Counter(r["type_reprise"] for r in uniques)

    out = args.out or str(BASE / "rapports" / f"retombees-{args.cp}.xlsx")
    pathlib.Path(out).parent.mkdir(parents=True, exist_ok=True)

    # ---- styles ----
    F_TITLE = Font(bold=True, size=14, color="FFFFFF")
    F_HDR = Font(bold=True, size=10, color="FFFFFF")
    F_CELL = Font(size=10)
    F_LINK = Font(size=10, color="0563C1", underline="single")
    F_BROUILLON = Font(bold=True, size=11, color="9A6B00")
    FILL_HDR = PatternFill("solid", fgColor=PRIM)
    FILL_BROUILLON = PatternFill("solid", fgColor="FDF3D1")
    FILL_ALT = PatternFill("solid", fgColor="F5F8FB")
    THIN = Side(style="thin", color="C9D2DC")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    def banner(ws, ncols):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        c = ws.cell(row=1, column=1, value=filigrane)
        c.font = F_BROUILLON
        c.fill = FILL_BROUILLON
        c.alignment = Alignment(horizontal="center")

    def header_row(ws, row, headers):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font = F_HDR
            c.fill = FILL_HDR
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)

    # ===== 1. EXECUTIVE SUMMARY =====
    ws = wb.active
    ws.title = "EXECUTIVE SUMMARY"
    banner(ws, 2)
    ws.cell(row=3, column=1, value="Communiqué").font = Font(bold=True)
    ws.cell(row=3, column=2, value=titre).alignment = WRAP
    kpis = [
        ("Reprises confirmées", nb_conf),
        ("Reprises probables", nb_prob),
        ("Publications non vérifiées", nb_nv),
        ("Médias uniques", medias_uniques),
        ("Pays identifiés", len(pays_counts)),
        ("Langues", "Français"),
    ]
    r = 5
    for k, v in kpis:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    ws.cell(row=r + 1, column=1, value="Répartition par type").font = Font(bold=True)
    rr = r + 2
    for t, v in type_counts.most_common():
        ws.cell(row=rr, column=1, value=t)
        ws.cell(row=rr, column=2, value=v)
        rr += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    # ===== 2. MEDIA COVERAGE =====
    ws2 = wb.create_sheet("MEDIA COVERAGE")
    headers = ["ID", "Média", "Pays", "Date", "Titre", "Statut", "URL"]
    banner(ws2, len(headers))
    header_row(ws2, 2, headers)
    row = 3
    for r in uniques:
        statut = "Confirmée" if r.get("classification") == "confirmee" else "Probable"
        vals = [r["id"], r["media"], r.get("pays") or "—",
                r.get("date_publication") or "n/d", r["titre"], statut]
        for i, v in enumerate(vals, 1):
            c = ws2.cell(row=row, column=i, value=v)
            c.font = F_CELL
            c.border = BORDER
            c.alignment = WRAP
        u = ws2.cell(row=row, column=7, value=r["url"])
        u.hyperlink = r["url"]
        u.font = F_LINK
        u.border = BORDER
        u.alignment = WRAP
        if row % 2 == 1:
            for i in range(1, 8):
                ws2.cell(row=row, column=i).fill = FILL_ALT
        row += 1
    for i, w in enumerate([22, 30, 16, 12, 55, 12, 55], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A3"

    # ===== 3. GEOGRAPHIC ANALYSIS =====
    ws3 = wb.create_sheet("GEOGRAPHIC ANALYSIS")
    banner(ws3, 2)
    header_row(ws3, 2, ["Pays", "Nombre de reprises"])
    row = 3
    for pays, v in pays_counts.most_common():
        ws3.cell(row=row, column=1, value=pays)
        ws3.cell(row=row, column=2, value=v)
        row += 1
    chart = BarChart()
    chart.type = "col"
    chart.title = "Reprises par pays"
    data = Reference(ws3, min_col=2, min_row=2, max_row=row - 1)
    cats = Reference(ws3, min_col=1, min_row=3, max_row=row - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    ws3.add_chart(chart, "D3")
    if pays_nd:
        ws3.cell(row=row + 1, column=1, value="Pays non déterminés : " +
                 ", ".join(r["media"] for r in pays_nd)).font = Font(italic=True)
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 20

    # ===== 4. PUBLICATION TIMELINE =====
    ws4 = wb.create_sheet("PUBLICATION TIMELINE")
    headers4 = ["Date", "ID", "Média", "Titre", "Statut"]
    banner(ws4, len(headers4))
    header_row(ws4, 2, headers4)
    tri = sorted(uniques, key=lambda r: r.get("date_publication") or "9999")
    row = 3
    for r in tri:
        statut = "Confirmée" if r.get("classification") == "confirmee" else "Probable"
        vals = [r.get("date_publication") or "n/d", r["id"], r["media"], r["titre"], statut]
        for i, v in enumerate(vals, 1):
            c = ws4.cell(row=row, column=i, value=v)
            c.font = F_CELL
            c.border = BORDER
            c.alignment = WRAP
        row += 1
    for i, w in enumerate([12, 22, 30, 60, 12], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = "A3"

    # ===== 5. VERIFICATION LOG =====
    ws5 = wb.create_sheet("VERIFICATION LOG")
    headers5 = ["ID", "Média", "Statut", "Classification", "Source de recherche", "Date de collecte", "URL", "Note"]
    banner(ws5, len(headers5))
    header_row(ws5, 2, headers5)
    row = 3
    for r in uniques + nv:
        vals = [r["id"], r["media"], r.get("statut_verification", ""),
                r.get("classification", ""), r.get("source_recherche", ""),
                r.get("date_collecte", ""), r["url"], r.get("note", "")]
        for i, v in enumerate(vals, 1):
            c = ws5.cell(row=row, column=i, value=v)
            c.font = F_CELL
            c.border = BORDER
            c.alignment = WRAP
        u = ws5.cell(row=row, column=7, value=r["url"])
        u.hyperlink = r["url"]
        u.font = F_LINK
        u.border = BORDER
        u.alignment = WRAP
        row += 1
    for i, w in enumerate([22, 30, 14, 16, 22, 14, 50, 55], 1):
        ws5.column_dimensions[get_column_letter(i)].width = w
    ws5.freeze_panes = "A3"

    wb.save(out)
    print(f"Excel généré : {out}")
    print(f"  onglets=5, confirmées={nb_conf}, probables={nb_prob}, non_vérifiées={nb_nv}")


if __name__ == "__main__":
    main()
