# -*- coding: utf-8 -*-
"""
Vérifie la cohérence entre reprises.csv et les fichiers PDF / XLSX générés.

Usage : python scripts/verifier_rapports.py [CP-XXXX]
"""
import csv
import pathlib
import sys
from collections import Counter

BASE = pathlib.Path(__file__).resolve().parent.parent


def verifier_csv():
    rows = list(csv.DictReader((BASE / "reprises.csv").open(encoding="utf-8"), delimiter=";"))
    reelles = [r for r in rows if r.get("est_test") == "non"]
    c = Counter(r.get("classification") for r in reelles)
    return {
        "nb_colonnes": len(rows[0]),
        "lignes_reelles": len(reelles),
        "classification": dict(c),
        "colonnes": rows[0],
    }


def verifier_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(str(path))
    txt = "".join((p.extract_text() or "") for p in r.pages)
    return {
        "pages": len(r.pages),
        "titre_ok": "conduite plus sûre" in txt,
        "executive_summary": "Executive Summary" in txt,
        "media_coverage": "Media Coverage" in txt,
        "geographic_coverage": "Geographic Coverage" in txt,
        "methodology": "Methodology" in txt,
        "brouillon": ("BROUILLON" in txt) or ("ATTENTE DE VALIDATION" in txt),
        "nb_caracteres": len(txt),
    }


def verifier_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    ws = wb["MEDIA COVERAGE"]
    n_lignes = 0
    n_liens = 0
    for row in ws.iter_rows(min_row=3, max_col=7):
        if row[0].value:
            n_lignes += 1
            if row[6].hyperlink:
                n_liens += 1
    return {
        "onglets": wb.sheetnames,
        "lignes_media_coverage": n_lignes,
        "liens_cliquables": n_liens,
        "brouillon_sheet1": wb["EXECUTIVE SUMMARY"].cell(row=1, column=1).value,
    }


def main():
    cp = sys.argv[1] if len(sys.argv) > 1 else "CP-20261010-001"
    print("=== CSV ===")
    print(verifier_csv())
    print("=== PDF ===")
    print(verifier_pdf(BASE / f"rapports/rapport-{cp}.pdf"))
    print("=== XLSX ===")
    print(verifier_xlsx(BASE / f"rapports/retombees-{cp}.xlsx"))


if __name__ == "__main__":
    main()
